from flask import Flask, request, jsonify, send_file
import pandas as pd
from geopy.distance import geodesic
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from io import BytesIO

app = Flask(__name__)

SECRET_KEY ="Rinben@123$$"

@app.route("/process", methods=["POST"])
def process_file():

    api_key = request.headers.get("x-api-key")
    if api_key != SECRET_KEY:
        return jsonify({"error": "Unauthorized"}), 401

    if "file" not in request.files:
        return jsonify({"error": "No file uploaded"}), 400

    file = request.files["file"]
    filename = file.filename

    df = pd.read_csv(file)
    df.columns = df.columns.str.strip()

    extracted_cols = df.columns[:5]
    df = df[extracted_cols]
    df.columns = ["ObjectId", "Latitude", "Longitude", "Altitude", "Name"]

    df = df.sort_values("ObjectId").reset_index(drop=True)

    df_temp = df.copy()
    df_temp["Sl_No"] = None

    current_index = 0
    order = 1
    df_temp.loc[current_index, "Sl_No"] = order

    while order < len(df_temp):
        current_point = (
            df_temp.loc[current_index, "Latitude"],
            df_temp.loc[current_index, "Longitude"]
        )

        min_dist = float("inf")
        next_index = None

        for i in range(len(df_temp)):
            if df_temp.loc[i, "Sl_No"] is None:
                candidate = (
                    df_temp.loc[i, "Latitude"],
                    df_temp.loc[i, "Longitude"]
                )
                d = geodesic(current_point, candidate).meters

                if d < min_dist:
                    min_dist = d
                    next_index = i

        order += 1
        df_temp.loc[next_index, "Sl_No"] = order
        current_index = next_index

    df["Sl_No"] = df_temp["Sl_No"]
    df = df.sort_values("Sl_No").reset_index(drop=True)

    distances = [0.00]
    for i in range(1, len(df)):
        prev = (df.loc[i - 1, "Latitude"], df.loc[i - 1, "Longitude"])
        curr = (df.loc[i, "Latitude"], df.loc[i, "Longitude"])
        d = geodesic(prev, curr).meters
        distances.append(round(d, 2))

    df["Distance [m]"] = distances

    # ---------- COLUMN ORDER (VERY IMPORTANT FIX) ----------
    df = df[[
        "ObjectId",
        "Sl_No",
        "Latitude",
        "Longitude",
        "Altitude",
        "Name",
        "Distance [m]"
    ]]

    # ---------- Add Remarks ----------
    df["Remarks"] = ""

    output = BytesIO()
    df.to_excel(output, index=False)
    output.seek(0)

    wb = load_workbook(output)
    ws = wb.active

    tab = Table(displayName="DGPS_Table", ref=f"A1:H{len(df) + 1}")
    style = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    tab.tableStyleInfo = style
    ws.add_table(tab)
    ws.row_dimensions[1].height = 30

    final_output = BytesIO()
    wb.save(final_output)
    final_output.seek(0)

    return send_file(
        final_output,
        download_name=filename.replace(".csv", "_CLEAN.xlsx"),
        as_attachment=True,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

if __name__ == "__main__":
    # This is only for local testing, NOT for Render.
    app.run(host="0.0.0.0", port=5000, debug=True)
